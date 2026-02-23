"""
Document loader module - supports loading multiple document formats
"""
import os
from typing import List, Dict, Optional
from pathlib import Path
import PyPDF2
from docx import Document
import openpyxl
import markdown
from bs4 import BeautifulSoup


class DocumentLoader:
    """Document loader supporting PDF, Word, Excel, Markdown, TXT and other formats"""
    
    def __init__(self):
        self.supported_formats = {
            '.pdf': self._load_pdf,
            '.docx': self._load_docx,
            '.xlsx': self._load_xlsx,
            '.xls': self._load_xlsx,
            '.md': self._load_markdown,
            '.txt': self._load_text,
            '.html': self._load_html,
        }
    
    def load_file(self, file_path: str, extract_metadata: bool = True) -> Dict[str, any]:
        """
        Load a single file
        
        Args:
            file_path: Path to the file
            extract_metadata: Whether to extract metadata (e.g., organization info)
            
        Returns:
            Dictionary containing document content and metadata
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        ext = file_path.suffix.lower()
        if ext not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {ext}")
        
        loader_func = self.supported_formats[ext]
        content = loader_func(file_path)
        
        # Extract language information (inferred from path)
        language = self._detect_language_from_path(str(file_path))
        
        # Extract topic information (inferred from path, suitable for macommunaute format)
        topic = self._extract_topic_from_path(str(file_path))
        
        metadata = {
            'size': file_path.stat().st_size,
            'modified_time': file_path.stat().st_mtime,
            'language': language,
            'topic': topic
        }
        
        # If macommunaute format, extract organization count
        if extract_metadata and ('macommunaute' in str(file_path).lower() or 
                                  '=====' in content[:500]):
            org_count = self._count_organizations(content)
            metadata['organization_count'] = org_count
        
        return {
            'file_path': str(file_path),
            'file_name': file_path.name,
            'file_type': ext,
            'content': content,
            'metadata': metadata
        }
    
    def _detect_language_from_path(self, file_path: str) -> str:
        """Detect language from file path"""
        file_path_lower = file_path.lower()
        if '/fr/' in file_path_lower or '_fr' in file_path_lower:
            return 'fr'
        elif '/en/' in file_path_lower or '_en' in file_path_lower:
            return 'en'
        elif '/zh/' in file_path_lower or '_zh' in file_path_lower:
            return 'zh'
        return 'unknown'
    
    def _extract_topic_from_path(self, file_path: str) -> str:
        """Extract topic from file path"""
        file_path_lower = file_path.lower()
        # Extract filename (without extension) as topic
        from pathlib import Path
        topic = Path(file_path).stem
        # Remove language suffix
        for lang in ['_fr', '_en', '_zh', '_fr_en']:
            if topic.endswith(lang):
                topic = topic[:-len(lang)]
        return topic
    
    def _count_organizations(self, content: str) -> int:
        """Count the number of organizations in the document"""
        import re
        pattern = r'=====\s*(?:Organisation|Organization)\s+\d+:\s*.+?\s*====='
        matches = re.findall(pattern, content)
        return len(matches)
    
    def load_directory(self, directory: str, recursive: bool = True) -> List[Dict[str, any]]:
        """
        Load all supported format files in a directory
        
        Args:
            directory: Directory path
            recursive: Whether to recursively load subdirectories
            
        Returns:
            List of documents
        """
        directory = Path(directory)
        if not directory.exists():
            raise FileNotFoundError(f"Directory not found: {directory}")
        
        documents = []
        pattern = '**/*' if recursive else '*'
        
        for file_path in directory.glob(pattern):
            if file_path.is_file() and file_path.suffix.lower() in self.supported_formats:
                try:
                    doc = self.load_file(file_path)
                    documents.append(doc)
                except Exception as e:
                    print(f"Failed to load file {file_path}: {e}")
        
        return documents
    
    def _load_pdf(self, file_path: Path) -> str:
        """Load PDF file"""
        content = []
        with open(file_path, 'rb') as f:
            pdf_reader = PyPDF2.PdfReader(f)
            for page in pdf_reader.pages:
                content.append(page.extract_text())
        return '\n\n'.join(content)
    
    def _load_docx(self, file_path: Path) -> str:
        """Load Word document"""
        doc = Document(file_path)
        content = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                content.append(paragraph.text)
        return '\n'.join(content)
    
    def _load_xlsx(self, file_path: Path) -> str:
        """Load Excel file"""
        workbook = openpyxl.load_workbook(file_path)
        content = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            content.append(f"Sheet: {sheet_name}")
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) if cell else '' for cell in row)
                if row_text.strip():
                    content.append(row_text)
            content.append('')
        return '\n'.join(content)
    
    def _load_markdown(self, file_path: Path) -> str:
        """Load Markdown file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def _load_text(self, file_path: Path) -> str:
        """Load plain text file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def _load_html(self, file_path: Path) -> str:
        """Load HTML file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            soup = BeautifulSoup(f.read(), 'html.parser')
            return soup.get_text()
